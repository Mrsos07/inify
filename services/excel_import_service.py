# -*- coding: utf-8 -*-
"""
Excel Import Service - خدمة استيراد العقارات من Excel
"""

import os
import logging
from io import BytesIO
from typing import List, Dict, Tuple
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.conf import settings

logger = logging.getLogger(__name__)

# الأعمدة المطلوبة في ملف Excel
EXCEL_COLUMNS = {
    'title': 'عنوان العقار *',
    'property_type': 'نوع العقار *',
    'status': 'حالة العرض *',
    'city': 'المدينة *',
    'neighborhood': 'الحي',
    'street': 'الشارع',
    'price': 'السعر *',
    'rent_period': 'فترة الإيجار',
    'size': 'المساحة (م²) *',
    'bedrooms': 'غرف النوم',
    'bathrooms': 'الحمامات',
    'living_rooms': 'غرف المعيشة',
    'floor_number': 'رقم الطابق',
    'parking_spaces': 'مواقف السيارات',
    'furnishing': 'حالة التأثيث',
    'year_built': 'سنة البناء',
    'description': 'الوصف',
    'is_negotiable': 'قابل للتفاوض',
    'amenities': 'المميزات',
}

# القيم المسموحة
PROPERTY_TYPES = {
    'شقة': 'apartment',
    'فيلا': 'villa',
    'مكتب': 'office',
    'أرض': 'land',
    'تجاري': 'commercial',
    'دوبلكس': 'duplex',
    'استوديو': 'studio',
    'مستودع': 'warehouse',
    'عمارة': 'building',
}

STATUS_TYPES = {
    'للبيع': 'for_sale',
    'للإيجار': 'for_rent',
    'محجوز': 'reserved',
    'مؤجر': 'rented',
    'مباع': 'sold',
}

FURNISHING_TYPES = {
    'مفروش': 'furnished',
    'نصف مفروش': 'semi_furnished',
    'غير مفروش': 'unfurnished',
}

RENT_PERIODS = {
    'سنوي': 'yearly',
    'شهري': 'monthly',
    'يومي': 'daily',
}

AMENITIES_MAP = {
    'مسبح': 'pool',
    'حديقة': 'garden',
    'صالة رياضية': 'gym',
    'حراسة': 'security',
    'مصعد': 'elevator',
    'موقف': 'parking',
    'تكييف مركزي': 'central_ac',
    'شرفة': 'balcony',
    'غرفة خادمة': 'maid_room',
    'غرفة سائق': 'driver_room',
    'مخزن': 'storage',
    'مطبخ': 'kitchen',
}


def generate_template() -> BytesIO:
    """
    إنشاء نموذج Excel للتحميل
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "العقارات"
    
    # تنسيق الهيدر
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # إضافة الهيدر
    headers = list(EXCEL_COLUMNS.values())
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        ws.column_dimensions[cell.column_letter].width = 18
    
    # إضافة صف مثال
    example_row = [
        'شقة فاخرة في الرياض',  # عنوان
        'شقة',  # نوع
        'للبيع',  # حالة
        'الرياض',  # مدينة
        'النرجس',  # حي
        'شارع الملك عبدالله',  # شارع
        '850000',  # سعر
        'سنوي',  # فترة الإيجار
        '180',  # مساحة
        '3',  # غرف نوم
        '2',  # حمامات
        '1',  # غرف معيشة
        '2',  # رقم الطابق
        '1',  # مواقف
        'غير مفروش',  # تأثيث
        '2020',  # سنة البناء
        'شقة فاخرة بتشطيبات عالية الجودة',  # وصف
        'نعم',  # قابل للتفاوض
        'مصعد، موقف، تكييف مركزي',  # مميزات
    ]
    
    example_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    for col, value in enumerate(example_row, 1):
        cell = ws.cell(row=2, column=col, value=value)
        cell.fill = example_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # إضافة ورقة التعليمات
    ws_help = wb.create_sheet("التعليمات")
    ws_help.sheet_view.rightToLeft = True
    
    instructions = [
        ("تعليمات استخدام ملف استيراد العقارات", ""),
        ("", ""),
        ("الحقول الإلزامية:", "عنوان العقار، نوع العقار، حالة العرض، المدينة، السعر، المساحة"),
        ("", ""),
        ("أنواع العقارات المتاحة:", "شقة، فيلا، مكتب، أرض، تجاري، دوبلكس، استوديو، مستودع، عمارة"),
        ("", ""),
        ("حالات العرض:", "للبيع، للإيجار، محجوز، مؤجر، مباع"),
        ("", ""),
        ("حالات التأثيث:", "مفروش، نصف مفروش، غير مفروش"),
        ("", ""),
        ("فترات الإيجار:", "سنوي، شهري، يومي (تُستخدم فقط للعقارات المعروضة للإيجار)"),
        ("", ""),
        ("قابل للتفاوض:", "نعم أو لا"),
        ("", ""),
        ("المميزات:", "اكتب المميزات مفصولة بفاصلة، مثال: مسبح، حديقة، مصعد، موقف"),
        ("", ""),
        ("المميزات المتاحة:", "مسبح، حديقة، صالة رياضية، حراسة، مصعد، موقف، تكييف مركزي، شرفة، غرفة خادمة، غرفة سائق، مخزن، مطبخ"),
        ("", ""),
        ("ملاحظات:", ""),
        ("- الصف الثاني في ورقة العقارات هو مثال، يمكنك حذفه أو التعديل عليه", ""),
        ("- تأكد من كتابة القيم بالضبط كما في التعليمات", ""),
        ("- الأرقام يجب أن تكون أرقام فقط بدون فواصل أو رموز", ""),
    ]
    
    title_font = Font(bold=True, size=14, color="4F46E5")
    subtitle_font = Font(bold=True, size=11)
    
    for row, (label, value) in enumerate(instructions, 1):
        cell_a = ws_help.cell(row=row, column=1, value=label)
        cell_b = ws_help.cell(row=row, column=2, value=value)
        
        if row == 1:
            cell_a.font = title_font
        elif label and not value:
            cell_a.font = subtitle_font
        
        ws_help.column_dimensions['A'].width = 30
        ws_help.column_dimensions['B'].width = 80
    
    # حفظ في BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return output


def parse_excel_file(file) -> Tuple[List[Dict], List[str]]:
    """
    قراءة ومعالجة ملف Excel
    
    Returns:
        Tuple[List[Dict], List[str]]: قائمة العقارات وقائمة الأخطاء
    """
    properties = []
    errors = []
    
    try:
        wb = load_workbook(file, data_only=True)
        ws = wb.active
        
        # قراءة الهيدر
        headers = []
        for cell in ws[1]:
            headers.append(cell.value)
        
        # التحقق من الأعمدة الإلزامية
        required_headers = ['عنوان العقار *', 'نوع العقار *', 'حالة العرض *', 'المدينة *', 'السعر *', 'المساحة (م²) *']
        missing_headers = [h for h in required_headers if h not in headers]
        
        if missing_headers:
            errors.append(f"أعمدة مفقودة: {', '.join(missing_headers)}")
            return properties, errors
        
        # إنشاء خريطة الأعمدة
        header_to_field = {v: k for k, v in EXCEL_COLUMNS.items()}
        col_map = {}
        for idx, header in enumerate(headers):
            if header in header_to_field:
                col_map[header_to_field[header]] = idx
        
        # قراءة الصفوف
        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            # تخطي الصفوف الفارغة
            if not any(row):
                continue
            
            property_data = {}
            row_errors = []
            
            # قراءة كل عمود
            for field, col_idx in col_map.items():
                value = row[col_idx] if col_idx < len(row) else None
                
                if value is not None:
                    value = str(value).strip()
                
                # معالجة القيم
                if field == 'property_type':
                    if value:
                        mapped = PROPERTY_TYPES.get(value)
                        if mapped:
                            property_data[field] = mapped
                        else:
                            row_errors.append(f"نوع العقار غير صحيح: {value}")
                    else:
                        row_errors.append("نوع العقار مطلوب")
                        
                elif field == 'status':
                    if value:
                        mapped = STATUS_TYPES.get(value)
                        if mapped:
                            property_data[field] = mapped
                        else:
                            row_errors.append(f"حالة العرض غير صحيحة: {value}")
                    else:
                        row_errors.append("حالة العرض مطلوبة")
                        
                elif field == 'furnishing':
                    if value:
                        mapped = FURNISHING_TYPES.get(value)
                        if mapped:
                            property_data[field] = mapped
                        else:
                            property_data[field] = 'unfurnished'
                    else:
                        property_data[field] = 'unfurnished'
                        
                elif field == 'rent_period':
                    if value:
                        mapped = RENT_PERIODS.get(value)
                        if mapped:
                            property_data[field] = mapped
                        else:
                            property_data[field] = 'yearly'
                    else:
                        property_data[field] = 'yearly'
                        
                elif field == 'is_negotiable':
                    property_data[field] = value in ['نعم', 'yes', 'true', '1', 'صح'] if value else True
                    
                elif field == 'amenities':
                    if value:
                        amenities = []
                        for item in value.split('،'):
                            item = item.strip()
                            if not item:
                                continue
                            # محاولة بالفاصلة الإنجليزية أيضاً
                            for sub_item in item.split(','):
                                sub_item = sub_item.strip()
                                if sub_item in AMENITIES_MAP:
                                    amenities.append(AMENITIES_MAP[sub_item])
                        property_data['amenities'] = amenities
                        
                elif field in ['price', 'size']:
                    if value:
                        try:
                            # إزالة الفواصل والعملات
                            clean_value = ''.join(c for c in str(value) if c.isdigit() or c == '.')
                            property_data[field] = float(clean_value)
                        except (ValueError, TypeError):
                            row_errors.append(f"{EXCEL_COLUMNS[field]} يجب أن يكون رقماً: {value}")
                    else:
                        row_errors.append(f"{EXCEL_COLUMNS[field]} مطلوب")
                        
                elif field in ['bedrooms', 'bathrooms', 'living_rooms', 'floor_number', 'parking_spaces', 'year_built']:
                    if value:
                        try:
                            property_data[field] = int(float(value))
                        except (ValueError, TypeError):
                            pass  # تجاهل الأخطاء للحقول الاختيارية
                            
                elif field in ['title', 'city']:
                    if value:
                        property_data[field] = value
                    else:
                        row_errors.append(f"{EXCEL_COLUMNS[field]} مطلوب")
                        
                else:
                    if value:
                        property_data[field] = value
            
            # إضافة العقار أو الأخطاء
            if row_errors:
                errors.append(f"صف {row_num}: " + "، ".join(row_errors))
            else:
                properties.append(property_data)
        
        if not properties and not errors:
            errors.append("الملف فارغ أو لا يحتوي على بيانات صحيحة")
            
    except Exception as e:
        logger.error(f"Error parsing Excel file: {e}", exc_info=True)
        errors.append(f"خطأ في قراءة الملف: {str(e)}")
    
    return properties, errors


def import_properties(agent, properties_data: List[Dict]) -> Tuple[int, List[str]]:
    """
    استيراد العقارات إلى قاعدة البيانات
    
    Args:
        agent: المسوق العقاري
        properties_data: قائمة بيانات العقارات
    
    Returns:
        Tuple[int, List[str]]: عدد العقارات المضافة وقائمة الأخطاء
    """
    from apps.properties.models import Property, PropertyAmenity
    
    imported = 0
    errors = []
    
    for idx, data in enumerate(properties_data, 1):
        try:
            # استخراج المميزات
            amenities = data.pop('amenities', [])
            
            # إنشاء العقار
            property_obj = Property.objects.create(
                agent=agent,
                **data
            )
            
            # إضافة المميزات
            for amenity in amenities:
                PropertyAmenity.objects.create(
                    property=property_obj,
                    amenity=amenity
                )
            
            imported += 1
            logger.info(f"Imported property: {property_obj.title}")
            
        except Exception as e:
            errors.append(f"عقار {idx}: {str(e)}")
            logger.error(f"Error importing property {idx}: {e}", exc_info=True)
    
    return imported, errors
